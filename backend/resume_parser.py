from io import BytesIO
from pypdf import PdfReader
from docx import Document

ALLOWED_EXTENSIONS = {"pdf", "docx"}

def extract_resume_text(file_storage):
    filename = (file_storage.filename or "").lower()
    if "." not in filename:
        raise ValueError("File must have a .pdf or .docx extension.")
    extension = filename.rsplit(".", 1)[1]
    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError("Unsupported file type. Please upload a PDF or DOCX file.")

    raw = file_storage.read()
    if not raw:
        raise ValueError("The uploaded file is empty.")
    try:
        if extension == "pdf":
            reader = PdfReader(BytesIO(raw))
            if reader.is_encrypted:
                raise ValueError("This PDF is password-protected. Upload an unlocked PDF.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            document = Document(BytesIO(raw))
            paragraphs = [p.text for p in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    paragraphs.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(paragraphs)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Could not read this file. Please try another PDF or DOCX.") from exc
    return text.strip()
